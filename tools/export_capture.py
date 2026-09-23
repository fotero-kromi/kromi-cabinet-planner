"""Capture the page's export artifacts headlessly and canonicalize them.

Drives the reload path on the real raw workbook, intercepts st.download_button
payloads, and writes a manifest of content digests with the volatile parts
neutralized: the Run_Metadata timestamp cell, the workbook's zip/package
timestamps (implicitly, by hashing cell values, not bytes), and the build
label (so a capture before a BUILD bump compares against one after it). Two
toggle configurations are captured: the default workbook and the technical
one. Excel-only since v34.25: the per-SP PDF was retired and its content
moved into the workbook.

Freezes the wall clock during the drive (freezegun, an instrument-only
dependency, not part of the app requirements) so the generated-at stamps in
the workbook metadata cannot vary between captures; the canonicalizer
additionally neutralizes the build label so a capture taken before a BUILD
bump compares clean against one taken after.

Usage: KROMI_RAW_INPUT_XLSX=<raw workbook> python3 tools/export_capture.py OUTDIR LABEL
Compare two labels by diffing OUTDIR/<LABEL>/manifest.json.
"""

import hashlib
import io
import json
import os
import re
import sys
import tempfile
from pathlib import Path

import pandas as pd

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO))
os.chdir(REPO)

# Drives must never write remembered KTC-IDs into the repository (v34.48).
os.environ.setdefault("KROMI_FILE_PREFS_PATH", str(Path(tempfile.mkdtemp()) / "file_prefs.json"))
from engine.run_restore import NOT_AVAIL, build_seed  # noqa: E402

RAW = os.environ["KROMI_RAW_INPUT_XLSX"]
OUT = Path(sys.argv[1]); LABEL = sys.argv[2]
(OUT / LABEL).mkdir(parents=True, exist_ok=True)

_BUILD_RE = re.compile(r"v34\.\d+")


def _canon_excel(xbytes: bytes) -> str:
    """Digest of every sheet's cell values, timestamp and build neutralized."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(xbytes), data_only=False)
    dump = {}
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            vals = []
            for v in row:
                s = "" if v is None else str(v)
                s = _BUILD_RE.sub("vX", s)
                vals.append(s)
            rows.append(vals)
        # Neutralize the Run_Metadata timestamp value (row: Key == 'Timestamp (UTC)')
        if ws.title == "Run_Metadata":
            rows = [["<ts>" if (i > 0 and r and r[0] == "Timestamp (UTC)" and j == 1)
                     else c for j, c in enumerate(r)] for i, r in enumerate(rows)]
        dump[ws.title] = rows
    blob = json.dumps(dump, sort_keys=True).encode()
    return hashlib.sha256(blob).hexdigest()


def _drive(controls: dict, technical: bool = False) -> dict:
    import streamlit as st
    from streamlit.testing.v1 import AppTest

    captured: dict = {}
    _orig_dl = st.download_button

    def spy_download(*args, **kw):
        try:
            label = kw.get("label", args[0] if args else "")
            data = kw.get("data", args[1] if len(args) > 1 else None)
            payload = data.getvalue() if hasattr(data, "getvalue") else data
            if isinstance(payload, (bytes, bytearray)):
                captured[str(label)] = bytes(payload)
        except Exception:
            pass
        return _orig_dl(*args, **kw)

    st.download_button = spy_download
    try:
        from freezegun import freeze_time
    except ImportError as exc:  # instrument-only dependency
        raise SystemExit(
            "export_capture needs freezegun (instrument-only): pip install freezegun"
        ) from exc
    try:
        raw_bytes = open(RAW, "rb").read()
        raw_cols = list(pd.ExcelFile(RAW).parse("Sheet1").columns)
        ui = {"cm_code": "WZIntNr", "cm_desc1": "WZBez",
              "cm_cons": "Consumption Last 16 Months",
              "cm_stdspecial": "WZArtID (1Standard/2Sonder)",
              "cm_sup": "WZLiefStamm.LWZBestellNr", "cm_site": "Werk",
              "ks_sheet_tools": "Sheet1",
              **controls}
        restore = build_seed(ui, raw_cols)
        restore["ks_hide_optional"] = False
        restore["ks_ai_colmap"] = False
        restore["_std_special_mapped"] = True
        with freeze_time("2026-07-05 12:00:00"), tempfile.TemporaryDirectory() as tmp:
            os.environ["KROMI_DB_PATH"] = str(Path(tmp) / "x.db")
            at = AppTest.from_file("pages/1_Kromi_Planner.py", default_timeout=300)
            at.session_state["_reload_ctx"] = {
                "run_id": 1, "bytes": raw_bytes, "filename": "input.xlsx",
                "customer": "TestCustomer", "site": "TestSite",
                "classifications": {}, "override_mode": "none", "override_set_id": None,
            }
            at.session_state["_pending_restore"] = restore
            at.run()
            assert not at.exception, f"page raised on pass 1: {at.exception}"
            # v34.56: the raw file has a stock column that is now found
            # automatically; the capture measures the planning exports, so the
            # takeover stock stays unmapped (as before v34.56). The new-file
            # reset clears a seeded optional pick, so it is selected here.
            _stock = next((sb for sb in at.selectbox if sb.key == "cm_stock"), None)
            if _stock is not None:
                _stock.select(NOT_AVAIL)
                at.run()
                assert not at.exception, f"page raised on pass 1b: {at.exception}"
            at.session_state["_force_run"] = True
            at.run()
            _btn = next((b for b in at.button
                         if getattr(b, "key", "") == "btn_prepare_exports"), None)
            if _btn is not None:
                _btn.click()
                at.run()
            assert not at.exception, f"page raised: {at.exception}"
            if technical:
                _cb = next(cb for cb in at.checkbox if "technical" in str(cb.label))
                _cb.set_value(True).run()
                assert not at.exception, f"page raised (technical): {at.exception}"

            excel = None
            for lbl, b in captured.items():
                if b[:2] == b"PK" and "plan workbook" in lbl.lower():
                    excel = b
    finally:
        st.download_button = _orig_dl
    out = {}
    _keep = os.environ.get("KROMI_EXPORT_KEEP")
    if _keep and excel is not None:
        Path(_keep).mkdir(parents=True, exist_ok=True)
        (Path(_keep) / f"{controls.get('ks_n_sp','1')}_{technical}.xlsx").write_bytes(bytes(excel))
    if excel is not None:
        out["excel_sheets_digest"] = _canon_excel(excel)
        out["excel_n_bytes_class"] = len(excel) // 4096
    return out


def main():
    manifest = {}
    for name, controls, tech in [("default_cfg", {}, False),
                           ("technical_cfg", {"ks_op_mode": "Helix + Carousel (capped)",
                                              "ks_n_sp": 2}, True)]:
        print(f"[{LABEL}] capturing {name} ...", flush=True)
        manifest[name] = _drive(controls, technical=tech)
        print(f"  {manifest[name]}")
    (OUT / LABEL / "manifest.json").write_text(json.dumps(manifest, indent=2, sort_keys=True))
    print(f"[{LABEL}] written")


main()
