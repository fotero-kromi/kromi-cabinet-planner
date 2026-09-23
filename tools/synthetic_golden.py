"""Synthetic golden gate (v34.59).

The golden and export capture tools need a private customer workbook, so the
byte-identity gates only ran by hand. This module generates an anonymous,
deterministic tool catalog, drives the planner page through four scenarios
(standard, capped with two replicated supply points, fixed configuration with
a Program mapping, stock and the stock-based Helix promotion, and the
numbering-only mode) and reduces each run to digests: the persisted plan
tables and the exported workbook. tests/test_synthetic_golden.py compares
them with the committed manifest in every test run, so a refactor that
changes a result fails in the suite and in CI.

Usage:
    python tools/synthetic_golden.py --check     compare with the manifest
    python tools/synthetic_golden.py --update    rewrite the manifest after an
                                                 intended change (say why in
                                                 the CHANGELOG)

The catalog is invented: no customer data, names or article numbers.
"""

from __future__ import annotations

import hashlib
import json
import os
import random
import sqlite3
import sys
import tempfile
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, List, Optional

import pandas as pd

_REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(_REPO_ROOT))

MANIFEST = _REPO_ROOT / "tests" / "golden" / "synthetic_manifest.json"
PAGE = str(_REPO_ROOT / "pages" / "1_Kromi_Planner.py")
SHEET = "Catalog"
SEED = 20260923

# ---- the catalog ------------------------------------------------------------------

# family: (weight, category words, description templates, pack units, consumption base)
_FAMILIES: Dict[str, tuple] = {
    "drill": (22, ["Spiralbohrer VHM", "Spiralbohrer HSS-E/HSCO", "NC-Anbohrer VHM", ""],
              ["VHM-Bohrer D{d:.1f} 5xD", "Drill HSS D{d:.1f} L{l}", "Spiralbohrer D{d:.1f} GL{l}"],
              [1, 1, 1, 5, 10], 60),
    "step": (2, ["Stufenbohrer VHM"], ["Drm. {d:.1f}/{d2:.1f}"], [1], 20),
    "mill": (20, ["Schaftfräser VHM", "Radiusfräser VHM", ""],
             ["Schaftfräser D{d:.0f} Z4", "End mill D{d:.0f} R0.5", "Torus mill D{d:.0f}"],
             [1], 40),
    "insert": (18, ["WSP Drehen VHM", "WSP Fräsen VHM", ""],
               ["CNMG 120408-PM", "WNMG 080408", "APKT 1604PDR", "SDHW 0903", "DCMT 11T304"],
               [10], 300),
    "tap": (8, ["Gewindebohrer HSS", ""], ["Gewindebohrer M{m}", "Tap M{m}x{p}"], [1], 30),
    "reamer": (5, ["Reibahle einstufig VHM"], ["Reibahle D{d:.0f} H7"], [1], 10),
    "holder": (4, [""], ["Schrumpffutter HSK63 D{d:.0f}", "Spannzangenfutter ER32 D{d:.0f}"],
               [1], 2),
    "screw": (5, ["Schraube", ""], ["Schraube M{m}x{l}", "Torx screw M{m}"], [100, 50], 900),
    "abrasive": (4, [""], ["Trennscheibe 125x1,0", "Schleifband 40x620 K80"], [25, 10], 400),
    "saw": (2, ["Kreissägeblatt HM"], ["Kreissägeblatt D{d:.0f}x2"], [1], 8),
    "other": (10, ["Freitext", ""], ["Sonderwerkzeug Nr {n}", "Messerkopf M{m}"], [1], 5),
}

COLUMNS = ["Article No", "Description", "Description 2", "Category",
           "Consumption 12 months", "Pack unit", "Location", "Stock", "System",
           "Supplier article no"]


def build_catalog(n: int = 360, seed: int = SEED) -> pd.DataFrame:
    """The deterministic synthetic tool catalog (``n`` rows plus duplicates)."""
    rng = random.Random(seed)
    names = list(_FAMILIES)
    weights = [_FAMILIES[f][0] for f in names]
    rows: List[Dict[str, Any]] = []
    for i in range(1, n + 1):
        fam = rng.choices(names, weights)[0]
        _w, cats, templates, packs, base = _FAMILIES[fam]
        d = rng.choice([1.5, 2.5, 3.0, 4.2, 5.0, 6.8, 8.0, 10.0, 12.0, 16.0, 20.0, 25.0, 32.0])
        desc = rng.choice(templates).format(
            d=d, d2=d * 1.6, l=rng.choice([40, 62, 86, 120]), m=rng.choice([3, 4, 5, 6, 8, 10, 12]),
            p=rng.choice([0.5, 0.7, 1.0, 1.25]), n=i)
        cons = 0 if rng.random() < 0.25 else min(5000, int(rng.paretovariate(1.3) * base * 0.4))
        stock = (rng.randint(0, 60) if cons == 0
                 else int(cons * rng.uniform(0.05, 0.6)) + rng.choice([0, 0, 2, 5, 10, 30]))
        system = rng.choices(["KTC", "", "KTC or Kanban", "Locker"], [50, 38, 10, 2])[0]
        if fam == "holder" and rng.random() < 0.4:
            system = "Locker"
        rows.append({
            "Article No": f"SYN-{i:05d}",
            "Description": desc,
            "Description 2": rng.choice(["", "", "beschichtet", "TiAlN", "Sonder"]),
            "Category": rng.choice(cats),
            "Consumption 12 months": cons,
            "Pack unit": rng.choice(packs),
            "Location": "Line A" if rng.random() < 0.6 else "Line B",
            "Stock": stock,
            "System": system,
            "Supplier article no": (desc.split()[0] + desc.split()[1] if fam == "insert"
                                    else f"S{rng.randint(1, 99999):05d}"),
        })
    # Repeated articles: the same code at the other location (kept apart by the
    # Program mapping) and at the same location (merged by the deduplication).
    for k in range(4):
        src = dict(rows[10 + k * 7])
        src["Location"] = "Line B" if src["Location"] == "Line A" else "Line A"
        rows.append(src)
        dup = dict(rows[40 + k * 5])
        dup["Consumption 12 months"] = int(dup["Consumption 12 months"]) + 7
        rows.append(dup)
    return pd.DataFrame(rows, columns=COLUMNS)


def catalog_bytes(frame: Optional[pd.DataFrame] = None) -> bytes:
    bio = BytesIO()
    with pd.ExcelWriter(bio, engine="openpyxl") as writer:
        (build_catalog() if frame is None else frame).to_excel(
            writer, index=False, sheet_name=SHEET)
    return bio.getvalue()


# ---- the scenarios --------------------------------------------------------------------

_REQUIRED = {"cm_code": "Article No", "cm_desc1": "Description",
             "cm_cons": "Consumption 12 months"}
_OPTIONAL_KEYS = ("cm_prod", "cm_year", "cm_desc2", "cm_program", "cm_restock", "cm_sup",
                  "cm_size", "cm_site", "cm_stdspecial", "cm_pack", "cm_dims", "cm_regrind",
                  "cm_systemtyp", "cm_stock")
_BASE_MAP = {"cm_prod": "Category", "cm_pack": "Pack unit", "cm_sup": "Supplier article no"}

SCENARIOS: Dict[str, Dict[str, Any]] = {
    "standard": {
        "controls": {"ks_empty_cab": 50.0, "ks_helix_threshold": 1.0},
        "mapping": {**_BASE_MAP, "cm_desc2": "Description 2"},
        "download": "plan workbook",
    },
    "capped_replicate": {
        "controls": {"ks_op_mode": "Helix + Carousel (capped)", "ks_helix_threshold": 2.0,
                     "ks_bulk_routing": True, "ks_force_screws": True, "ks_buffer": 30.0,
                     "ks_n_sp": 2, "ks_min_carousel": 6, "ks_max_carousels": 1},
        "mapping": dict(_BASE_MAP),
        "download": "plan workbook",
    },
    "fixed_program_stock": {
        "controls": {"ks_op_mode": "Fixed configuration (existing machines)", "ks_n_sp": 2,
                     "ks_consumption_months": 12.0, "ks_helix_threshold": 6.0,
                     "ks_fixed_headroom": 10.0,
                     "ks_fixed_helix_1": 1, "ks_fixed_carousel_1": 1,
                     "ks_fixed_helix_2": 1, "ks_fixed_carousel_2": 0,
                     "ks_fixed_locker_a_2": 1,
                     "ks_fixed_stock_promo": True, "ks_fixed_stock_months": 3.0},
        "mapping": {**_BASE_MAP, "cm_program": "Location", "cm_stock": "Stock",
                    "cm_systemtyp": "System"},
        "programs": {"Line A": 1, "Line B": 2},
        "download": "plan workbook",
    },
    "numbering_only": {
        "controls": {"ks_op_mode": "Only article number assignment"},
        "mapping": dict(_BASE_MAP),
        "download": "article setup",
    },
}

# ---- digests --------------------------------------------------------------------------

_PLAN_SQL = {
    "pertool": """
        SELECT t.line_no, t.code, t.listing, t.raw_consumption, t.pack_units, t.system_typ,
               c.supply_point, c.bucket_label, c.system_category, c.cabinet_type,
               c.size_category, c.product_category, c.spirals_needed, c.carousel_stockpiles,
               c.spiral_capacity, c.monthly_packs, c.target_packs, c.consumption_pcs,
               c.size_issue, c.size_issue_reason, c.override_applied, c.restockable,
               c.restock_slots
        FROM cabinet_calculations c
        JOIN tool_records t ON c.tool_record_id = t.tool_record_id
        ORDER BY t.line_no, t.code, c.supply_point, c.bucket_label, c.calc_id""",
    "buckets": """
        SELECT bucket_label, helix_cabs, carousel_cabs, locker_a, locker_b, locker_c,
               total_cabs, total_spirals, carousel_slots, ktc_count, kanban_count
        FROM cabinet_plan_summary ORDER BY bucket_label, summary_id""",
    "grand": """
        SELECT rows_processed, rebalance_moves, cabinets_saved, size_issues_found,
               grand_total_cabs
        FROM engine_executions ORDER BY execution_id""",
}
# Run_Metadata rows that describe the environment, not the plan.
_VOLATILE_META = {"Build", "Timestamp (UTC)", "Model"}


def _sha(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def frame_digest(frame: pd.DataFrame) -> str:
    """Content digest of a plan table, as CSV without the index.

    The line terminator is fixed: to_csv() defaults to os.linesep, which made
    every plan digest differ on Windows (v34.60)."""
    return _sha(frame.to_csv(index=False, lineterminator="\n"))


def workbook_digest(data: bytes) -> str:
    """Content digest of an exported workbook: every sheet, every cell as
    text, without the build, the timestamp and the AI model rows."""
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(data), read_only=True, data_only=False)
    dump: Dict[str, List[List[str]]] = {}
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            vals = ["" if v is None else str(v) for v in row]
            if ws.title == "Run_Metadata" and vals and vals[0] in _VOLATILE_META:
                continue
            rows.append(vals)
        dump[ws.title] = rows
    return _sha(json.dumps(dump, sort_keys=True))


def _plan_digests(db_path: str) -> Dict[str, Any]:
    out: Dict[str, Any] = {}
    if not os.path.exists(db_path):
        return out
    conn = sqlite3.connect(db_path)
    try:
        for name, sql in _PLAN_SQL.items():
            frame = pd.read_sql_query(sql, conn)
            out[name] = {"rows": int(len(frame)), "sha256": frame_digest(frame)}
    finally:
        conn.close()
    return out


# ---- the drive -----------------------------------------------------------------------------

def run_scenario(name: str, work_dir: str) -> Dict[str, Any]:
    """Drive the page through one scenario; return its digests."""
    import streamlit as st
    from streamlit.testing.v1 import AppTest

    from engine.run_restore import NOT_AVAIL, build_seed

    spec = SCENARIOS[name]
    db_path = str(Path(work_dir) / f"{name}.db")
    os.environ["KROMI_DB_PATH"] = db_path
    os.environ.setdefault("KROMI_FILE_PREFS_PATH", str(Path(work_dir) / "file_prefs.json"))
    frame = build_catalog()
    raw = catalog_bytes(frame)
    cols = list(frame.columns)

    captured: Dict[str, bytes] = {}
    orig_dl = st.download_button

    def spy(*args, **kw):
        label = str(kw.get("label", args[0] if args else ""))
        data = kw.get("data", args[1] if len(args) > 1 else None)
        payload = data.getvalue() if hasattr(data, "getvalue") else data
        if isinstance(payload, (bytes, bytearray)):
            captured[label] = bytes(payload)
        return orig_dl(*args, **kw)

    st.download_button = spy
    st.cache_data.clear()
    try:
        seed = build_seed({**_REQUIRED, **spec["controls"], "ks_sheet_tools": SHEET}, cols)
        seed.update({"ks_hide_optional": False, "ks_ai_colmap": False,
                     "_std_special_mapped": False})
        at = AppTest.from_file(PAGE, default_timeout=300)
        at.session_state["_reload_ctx"] = {
            "run_id": 0, "bytes": raw, "filename": "synthetic_catalog.xlsx",
            "customer": "Synthetic", "site": "Golden", "classifications": {},
            "override_mode": "none", "override_set_id": None}
        at.session_state["_pending_restore"] = seed
        at.run()
        _raise(at, name, "first pass")
        # Every optional mapping is set explicitly, so auto-detection changes
        # never shift the gate.
        present = {sb.key for sb in at.selectbox}
        for key in _OPTIONAL_KEYS:
            if key in present:
                at.selectbox(key=key).select(spec["mapping"].get(key, NOT_AVAIL))
        at.run()
        _raise(at, name, "mapping")
        for prog, sp in spec.get("programs", {}).items():
            at.selectbox(key=f"prog_sp::{prog}").select(sp)
        if spec.get("programs"):
            at.run()
            _raise(at, name, "program mapping")
        at.session_state["_force_run"] = True
        at.run()
        _raise(at, name, "run")
        if spec["download"] == "plan workbook":
            btn = next(b for b in at.button if getattr(b, "key", "") == "btn_prepare_exports")
            btn.click()
            at.run()
            _raise(at, name, "exports")
    finally:
        st.download_button = orig_dl
    digests: Dict[str, Any] = dict(_plan_digests(db_path))
    book = next((b for lbl, b in captured.items()
                 if spec["download"] in lbl.lower() and b[:2] == b"PK"), None)
    if book is None:
        raise RuntimeError(f"{name}: the {spec['download']} download was not produced")
    digests["workbook"] = {"sha256": workbook_digest(book)}
    return digests


def _raise(at: Any, name: str, stage: str) -> None:
    if at.exception:
        raise RuntimeError(f"{name}: the page raised at the {stage}: {at.exception}")


def run_all() -> Dict[str, Any]:
    out: Dict[str, Any] = {}
    with tempfile.TemporaryDirectory() as tmp:
        for name in SCENARIOS:
            out[name] = run_scenario(name, tmp)
    return out


def load_manifest() -> Dict[str, Any]:
    return json.loads(MANIFEST.read_text(encoding="utf-8"))


def main(argv: List[str]) -> int:
    os.chdir(_REPO_ROOT)
    os.environ["KROMI_LOG_PATH"] = "off"
    os.environ["OPENAI_API_KEY"] = ""
    mode = argv[1] if len(argv) > 1 else "--check"
    got = run_all()
    if mode == "--update":
        MANIFEST.parent.mkdir(parents=True, exist_ok=True)
        MANIFEST.write_text(json.dumps({"scenarios": got}, indent=2, sort_keys=True) + "\n",
                            encoding="utf-8")
        print(f"manifest written: {MANIFEST}")
        return 0
    want = load_manifest()["scenarios"]
    bad = [f"{s}/{d}" for s in want for d in want[s] if want[s][d] != got.get(s, {}).get(d)]
    print("identical" if not bad else "DIFFERENT: " + ", ".join(bad))
    return 0 if not bad else 1


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
