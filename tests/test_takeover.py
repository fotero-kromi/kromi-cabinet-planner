"""Takeover sheets (v34.56): one sheet per supply point.

Every taken-over article is customer property. The sheet lists the
customer-property KROMI number, the customer number, both description texts,
the packaging unit, and how the customer's stock splits: whole packs go into
the KTC up to the article's maximum (the compartments allocated to it times
the packaging unit), the rest stays at the main stock location (HLO).

* Carousel maximum: compartments x VPE.
* Helix maximum: spirals x the places a spiral of the article's size holds
  (S 28, M 22, L 18, XL 12, category fallback) x VPE.
* Lockers, Kanban articles and articles without space: nothing in the KTC.
"""
import math

import pandas as pd
import pytest

from engine.fixed_config import STATUS_MOVED, STATUS_NOT_PLACED, STATUS_PLACED
from engine.takeover import (
    CATEGORY_TEXT_COL,
    STOCK_COL,
    TAKEOVER_COLUMNS,
    build_takeover_frames,
    build_takeover_workbook,
    guess_stock_column,
    max_packs,
    takeover_summary,
)


def _row(code, **kw):
    base = {
        "Listing": "Tools", "Code": code, "SupplyPoint": 1,
        "SystemCategory": "KTC", "CabinetType": "Carousel",
        "Spirals_needed": 0, "Spiral_capacity": pd.NA, "Carousel_stockpiles": 3,
        "PackUnits": 1.0, "SizeCategory": "M", "ProductCategory": "drills",
        "Placement_Status": STATUS_PLACED, "Regrind": False,
        "Description": f"desc {code}", "Description_2": "",
        "Kromi_Art_No": f"1401000{code[-2:]}000", STOCK_COL: 0.0,
    }
    base.update(kw)
    return base


def _frame(*rows):
    return pd.DataFrame(list(rows))


def _only(frames, sp=1):
    got = dict(frames)
    return got[sp].reset_index(drop=True)


# ---- the maximum per article ----------------------------------------------------------

@pytest.mark.parametrize("size,cap", [("S", 28), ("M", 22), ("L", 18), ("XL", 12)])
def test_helix_maximum_follows_the_spiral_size(size, cap):
    row = pd.Series(_row("A01", CabinetType="Helix", Spirals_needed=2,
                         Spiral_capacity=pd.NA, Carousel_stockpiles=0, SizeCategory=size))
    assert max_packs(row) == 2 * cap


def test_helix_maximum_uses_the_planned_spiral_capacity():
    row = pd.Series(_row("A01", CabinetType="Helix", Spirals_needed=1,
                         Spiral_capacity=22, SizeCategory="S"))
    assert max_packs(row) == 22


def test_helix_without_size_falls_back_by_category():
    ins = pd.Series(_row("A01", CabinetType="Helix", Spirals_needed=1, SizeCategory="",
                         ProductCategory="inserts"))
    mill = pd.Series(_row("A02", CabinetType="Helix", Spirals_needed=1, SizeCategory="",
                          ProductCategory="mills"))
    assert max_packs(ins) == 28 and max_packs(mill) == 18


def test_carousel_maximum_is_the_compartments():
    assert max_packs(pd.Series(_row("A01", Carousel_stockpiles=4))) == 4


def test_a_reground_helix_article_keeps_one_spiral_for_reground_pieces():
    row = pd.Series(_row("A01", CabinetType="Helix", Spirals_needed=2,
                         Spiral_capacity=22, Regrind=True))
    assert max_packs(row) == 22
    one = pd.Series(_row("A02", CabinetType="Helix", Spirals_needed=1,
                         Spiral_capacity=22, Regrind=True))
    assert max_packs(one) == 22


@pytest.mark.parametrize("kw", [
    {"CabinetType": "Locker A"},
    {"SystemCategory": "Kanban", "CabinetType": ""},
    {"Placement_Status": STATUS_NOT_PLACED},
])
def test_no_room_in_the_ktc(kw):
    assert max_packs(pd.Series(_row("A01", **kw))) == 0


def test_a_moved_article_uses_the_machine_it_moved_to():
    row = pd.Series(_row("A01", CabinetType="Helix", Spirals_needed=1, Spiral_capacity=18,
                         Carousel_stockpiles=0, Placement_Status=STATUS_MOVED))
    assert max_packs(row) == 18


# ---- the split ------------------------------------------------------------------------

def test_stock_fills_the_ktc_up_to_the_maximum_and_the_rest_goes_to_hlo():
    df = _frame(_row("A01", Carousel_stockpiles=3, PackUnits=10.0, **{STOCK_COL: 45.0}),
                _row("A02", Carousel_stockpiles=3, PackUnits=1.0, **{STOCK_COL: 2.0}))
    t = _only(build_takeover_frames(df))
    assert list(t["im KTC"]) == [30, 2]
    assert list(t["am HLO"]) == [15, 0]
    assert list(t["Total"]) == [45, 2]


def test_only_whole_packs_go_into_the_ktc():
    df = _frame(_row("A01", Carousel_stockpiles=5, PackUnits=10.0, **{STOCK_COL: 25.0}))
    t = _only(build_takeover_frames(df))
    assert (t.loc[0, "im KTC"], t.loc[0, "am HLO"], t.loc[0, "Total"]) == (20, 5, 25)


def test_helix_article_fills_its_spirals():
    df = _frame(_row("A01", CabinetType="Helix", Spirals_needed=1, Spiral_capacity=22,
                     Carousel_stockpiles=0, PackUnits=10.0, **{STOCK_COL: 300.0}))
    t = _only(build_takeover_frames(df))
    assert (t.loc[0, "im KTC"], t.loc[0, "am HLO"]) == (220, 80)


@pytest.mark.parametrize("kw", [
    {"CabinetType": "Locker B"},
    {"SystemCategory": "Kanban", "CabinetType": ""},
    {"Placement_Status": STATUS_NOT_PLACED},
])
def test_everything_stays_at_hlo_without_ktc_room(kw):
    df = _frame(_row("A01", **{STOCK_COL: 12.0}, **kw))
    t = _only(build_takeover_frames(df))
    assert (t.loc[0, "im KTC"], t.loc[0, "am HLO"], t.loc[0, "Total"]) == (0, 12, 12)


def test_missing_or_negative_stock_counts_as_zero_and_vpe_defaults_to_one():
    df = _frame(_row("A01", PackUnits=float("nan"), **{STOCK_COL: float("nan")}),
                _row("A02", PackUnits=0.0, **{STOCK_COL: -4.0}),
                _row("A03", PackUnits=float("nan"), **{STOCK_COL: 2.0}))
    t = _only(build_takeover_frames(df))
    assert list(t["Total"]) == [0, 0, 2]
    assert list(t["VPE"]) == [1, 1, 1]
    assert list(t["im KTC"]) == [0, 0, 2]


# ---- one sheet per supply point ------------------------------------------------------

def test_one_frame_per_supply_point_in_order_with_the_template_columns():
    df = _frame(_row("A01", SupplyPoint=2, **{STOCK_COL: 1.0}),
                _row("A02", SupplyPoint=1, **{STOCK_COL: 1.0}),
                _row("A03", SupplyPoint=2, **{STOCK_COL: 1.0}))
    frames = build_takeover_frames(df)
    assert [sp for sp, _ in frames] == [1, 2]
    for _, f in frames:
        assert list(f.columns) == list(TAKEOVER_COLUMNS)
    assert list(dict(frames)[2]["Kunden Art. Nr."]) == ["A01", "A03"]


def test_columns_follow_the_takeover_template():
    assert TAKEOVER_COLUMNS == (
        "Cust. Prop. Art. Nr.", "Kunden Art. Nr.", "Bezeichnung 1", "Bezeichnung 2",
        "VPE", "im KTC", "am HLO", "Total", "Schranktyp")


def test_each_supply_point_takes_over_its_own_stock():
    """Program mapping: the same article listed for two locations carries
    each location's own stock."""
    df = _frame(_row("A01", SupplyPoint=1, Carousel_stockpiles=3, **{STOCK_COL: 10.0}),
                _row("A01", SupplyPoint=2, Carousel_stockpiles=3, **{STOCK_COL: 1.0}))
    frames = dict(build_takeover_frames(df))
    assert (frames[1].loc[0, "im KTC"], frames[1].loc[0, "am HLO"]) == (3, 7)
    assert (frames[2].loc[0, "im KTC"], frames[2].loc[0, "am HLO"]) == (1, 0)


def test_replicated_articles_share_one_stock():
    """Replicate mode copies each article to every supply point with the same
    stock value: the stock fills the machines in supply-point order and the
    rest is listed once, on the first supply point's sheet."""
    df = _frame(_row("A01", SupplyPoint=1, Carousel_stockpiles=3, **{STOCK_COL: 10.0}),
                _row("A01", SupplyPoint=2, Carousel_stockpiles=3, **{STOCK_COL: 10.0}))
    frames = dict(build_takeover_frames(df, shared_stock=True))
    assert (frames[1].loc[0, "im KTC"], frames[1].loc[0, "am HLO"]) == (3, 4)
    assert (frames[2].loc[0, "im KTC"], frames[2].loc[0, "am HLO"]) == (3, 0)
    assert frames[1].loc[0, "Total"] + frames[2].loc[0, "Total"] == 10


def test_duplicate_rows_at_one_supply_point_become_one_line():
    df = _frame(_row("A01", Carousel_stockpiles=3, **{STOCK_COL: 2.0}),
                _row("A01", Carousel_stockpiles=4, **{STOCK_COL: 6.0}))
    t = _only(build_takeover_frames(df))
    assert len(t) == 1
    assert (t.loc[0, "im KTC"], t.loc[0, "am HLO"], t.loc[0, "Total"]) == (7, 1, 8)


def test_the_number_is_the_first_occurrence_customer_property_number():
    df = _frame(_row("A01", SupplyPoint=1, Kromi_Art_No="140100001000", **{STOCK_COL: 1.0}),
                _row("A01", SupplyPoint=2, Kromi_Art_No="140100002000", **{STOCK_COL: 1.0}))
    frames = dict(build_takeover_frames(df, shared_stock=True))
    assert frames[1].loc[0, "Cust. Prop. Art. Nr."] == "140100001000"
    assert frames[2].loc[0, "Cust. Prop. Art. Nr."] == "140100001000"


def test_numbers_stay_blank_when_they_could_not_be_generated():
    df = _frame(_row("A01", **{STOCK_COL: 1.0})).drop(columns=["Kromi_Art_No"])
    assert _only(build_takeover_frames(df)).loc[0, "Cust. Prop. Art. Nr."] == ""


def test_description_texts_follow_the_kds_template():
    kds = _frame(_row("A01", Description="Drm. 4,80/ 8,00", **{
        CATEGORY_TEXT_COL: "Stufenbohrer VHM", STOCK_COL: 1.0}))
    t = _only(build_takeover_frames(kds))
    assert (t.loc[0, "Bezeichnung 1"], t.loc[0, "Bezeichnung 2"]) == (
        "Stufenbohrer VHM", "Drm. 4,80/ 8,00")
    plain = _frame(_row("A01", Description="Drill D5", Description_2="long text",
                        **{STOCK_COL: 1.0}))
    t = _only(build_takeover_frames(plain))
    assert (t.loc[0, "Bezeichnung 1"], t.loc[0, "Bezeichnung 2"]) == ("Drill D5", "long text")


def test_machine_column_names_where_the_article_goes():
    df = _frame(_row("A01", CabinetType="Helix", Spirals_needed=1, **{STOCK_COL: 1.0}),
                _row("A02", **{STOCK_COL: 1.0}),
                _row("A03", Placement_Status=STATUS_NOT_PLACED, **{STOCK_COL: 1.0}),
                _row("A04", SystemCategory="Kanban", CabinetType="", **{STOCK_COL: 1.0}),
                _row("A05", CabinetType="Locker C", **{STOCK_COL: 1.0}))
    assert list(_only(build_takeover_frames(df))["Schranktyp"]) == [
        "Helix", "Carousel", "Not placed", "Kanban", "Locker C"]


def test_no_stock_column_means_no_takeover_sheets():
    df = _frame(_row("A01")).drop(columns=[STOCK_COL])
    assert build_takeover_frames(df) == []


def test_totals_are_conserved():
    df = _frame(*[_row(f"A{i:02d}", Carousel_stockpiles=i % 4 + 1, PackUnits=float(i % 3 + 1),
                       SupplyPoint=i % 2 + 1, **{STOCK_COL: float(i * 7 % 23)})
                  for i in range(1, 30)])
    frames = build_takeover_frames(df)
    total = sum(float(f["Total"].sum()) for _, f in frames)
    assert math.isclose(total, float(df[STOCK_COL].sum()))
    for _, f in frames:
        assert ((f["im KTC"] + f["am HLO"]) == f["Total"]).all()
        assert (f["im KTC"] % f["VPE"] == 0).all()


def test_summary_per_supply_point():
    df = _frame(_row("A01", SupplyPoint=1, **{STOCK_COL: 5.0}),
                _row("A02", SupplyPoint=2, PackUnits=10.0, **{STOCK_COL: 45.0}))
    s = takeover_summary(build_takeover_frames(df))
    assert s == [
        {"sp": 1, "articles": 1, "ktc": 3, "hlo": 2, "total": 5},
        {"sp": 2, "articles": 1, "ktc": 30, "hlo": 15, "total": 45},
    ]


def test_takeover_workbook_has_one_sheet_per_supply_point():
    from io import BytesIO
    from openpyxl import load_workbook
    df = _frame(_row("A01", SupplyPoint=1, **{STOCK_COL: 5.0}),
                _row("=HYPERLINK(1)", SupplyPoint=2, **{STOCK_COL: 5.0}))
    wb = load_workbook(BytesIO(build_takeover_workbook(build_takeover_frames(df))))
    assert wb.sheetnames == ["Takeover sheet SP 1", "Takeover sheet SP 2"]
    header = [c.value for c in wb["Takeover sheet SP 1"][1]]
    assert header == list(TAKEOVER_COLUMNS)
    cell = wb["Takeover sheet SP 2"]["B2"]
    assert cell.data_type != "f"          # customer text never becomes a formula


# ---- finding the stock column -----------------------------------------------------------

def test_stock_column_is_found_in_the_onboarding_header():
    cols = ["Kunden Artikelnummer / Customer Art.-No", "System ",
            "KTC Lagerorte / KTC locations", "AKTUELLER BESTAND 31.08.2026",
            "Preis- basis Stück / Price basis pcs.", "Verbrauch 12 Monate"]
    assert guess_stock_column(pd.DataFrame(columns=cols)) == "AKTUELLER BESTAND 31.08.2026"


def test_minimum_and_maximum_stock_columns_are_never_taken():
    cols = ["Kunden Artikelnummer / Customer Art.-No", "Mindestbestand / Min. stock",
            "Maximalbestand / Max. stock", "Lagerort / Stock location (Name)",
            "Jahresverbrauch / Yearly consumption"]
    assert guess_stock_column(pd.DataFrame(columns=cols)) is None


@pytest.mark.parametrize("header", ["Bestand", "Lagerbestand", "Stock", "Current stock",
                                    "Stock on hand", "Ist-Bestand"])
def test_common_stock_headers(header):
    assert guess_stock_column(pd.DataFrame(columns=["Code", header])) == header
