"""Contracts against silent failures (v34.48, audit C6 and C10).

C6: when the override database could not be read, the run proceeded without
any technician overrides and said "No stored override set for this scope
yet"; the technician panel then crashed before the save step. The run now
proceeds but tells the user plainly that overrides could not be loaded, and
the panel degrades to a message instead of crashing.

C10: KROMI numbers (and the Article setup sheet) were silently left out of the
plan workbook when the KTC-ID was invalid or a number field overflowed, while
the page still reported a clean export. The reason is now reported with the
export, and an invalid KTC-ID is flagged where it is typed, in every mode.
"""

from io import BytesIO

import pandas as pd

from engine.kromi_numbering import numbering_omission_reason
from engine.run_restore import build_seed
from tests._paths import PLANNER_PAGE


# ---- C10: numbering omission reasons ---------------------------------------

def _frame(n=1, desc="Bohrer D8,5", system="KTC"):
    return pd.DataFrame([{
        "Code": f"C{i}", "Listing": "Tools", "Description": desc,
        "SystemCategory": system, "CabinetType": "Helix",
        "ToolClass": "solid_carbide_drill", "PackUnits": 1.0,
    } for i in range(n)])


def test_no_reason_when_numbering_works():
    assert numbering_omission_reason(_frame(), "191") is None


def test_reason_for_invalid_ktc_id():
    for bad in ("", "19", "1911", "AB1"):
        r = numbering_omission_reason(_frame(), bad)
        assert r and "KTC-ID" in r, bad


def test_reason_for_missing_inputs():
    assert "ToolClass" in numbering_omission_reason(
        _frame().drop(columns=["ToolClass"]), "191")
    assert "description" in numbering_omission_reason(
        _frame().drop(columns=["Description"]), "191").lower()


def test_reason_for_variant_overflow():
    r = numbering_omission_reason(_frame(n=101, system="Kanban"), "191")
    assert r and "99" in r


def _workbook(ktc_id):
    from engine.workbook import build_result_workbook

    work = _frame()
    work["Monthly_packs"] = 1.0
    work["Spirals_needed"] = 1
    work["Carousel_stockpiles"] = 0
    empty = pd.DataFrame()
    ov = {"rows_touched": 0, "unmatched_overrides": [],
          "invalid_overrides": [], "applied_overrides": []}
    return build_result_workbook(
        work=work, df_summary=empty, df_audit=empty, df_bucket_compare=empty,
        presentation_compact_df=empty, presentation_detail_df=empty,
        dist_cat_rows=empty, dist_cat_vol=empty, dist_cabtype=empty,
        dist_system=empty, include_planogram=False, include_technical=False,
        ktc_id=ktc_id, apply_overrides_ui=False, enable_bulk_routing=False,
        multiple_listings=False, consumption_period_months=12.0,
        content_key="t", _df_run_meta=empty, _bucket_plans=[],
        _listings_arg=None, _listings_in_data=["Tools"], _base_info=None,
        _vend_stats={}, _override_stats=ov,
    )


def test_workbook_reports_why_numbers_are_missing():
    _x, problems, _p, _c, notes = _workbook("")
    assert "KROMI article numbers were not generated" in notes
    assert "KTC-ID" in notes
    assert problems == []          # a note, not a failed verification


def test_workbook_has_no_numbering_note_when_numbers_exist():
    _x, _pr, _p, _c, notes = _workbook("191")
    assert "KROMI article numbers" not in (notes or "")


# ---- page drives ------------------------------------------------------------

def _synthetic_workbook_bytes():
    df = pd.DataFrame({
        "ItemCode": [f"V{i:03d}" for i in range(6)],
        "ItemName": [f"Bohrer D{i + 3},5" for i in range(6)],
        "Cons16": [(i + 1) * 60.0 for i in range(6)],
    })
    bio = BytesIO()
    with pd.ExcelWriter(bio, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Sheet1")
    return bio.getvalue(), list(df.columns)


def _drive(monkeypatch, db_path, ktc_id=None):
    import streamlit as st
    from streamlit.testing.v1 import AppTest

    monkeypatch.setenv("KROMI_DB_PATH", str(db_path))
    st.cache_data.clear()
    raw, cols = _synthetic_workbook_bytes()
    restore = build_seed({"cm_code": "ItemCode", "cm_desc1": "ItemName",
                          "cm_cons": "Cons16", "ks_sheet_tools": "Sheet1"}, cols)
    restore["ks_hide_optional"] = False
    restore["ks_ai_colmap"] = False
    restore["_std_special_mapped"] = False
    at = AppTest.from_file(PLANNER_PAGE, default_timeout=300)
    at.session_state["_reload_ctx"] = {
        "run_id": 1, "bytes": raw, "filename": "input.xlsx",
        "customer": "TestCustomer", "site": "TestSite",
        "classifications": {}, "override_set_id": None}
    at.session_state["_pending_restore"] = restore
    at.run()
    if ktc_id is not None:
        next(t for t in at.text_input if t.label == "KTC-ID").input(ktc_id)
        at.run()
    at.session_state["_force_run"] = True
    at.run()
    return at


def test_unreadable_override_database_is_reported_not_swallowed(monkeypatch, tmp_path):
    bad = tmp_path / "corrupt.db"
    bad.write_bytes(b"this is not a sqlite database" * 200)
    at = _drive(monkeypatch, bad)
    assert not at.exception, at.exception
    warnings = " ".join(str(w.value) for w in at.warning)
    assert "override database could not be read" in warnings.lower()
    captions = " ".join(str(c.value) for c in at.caption)
    assert "Overrides applied from database set" not in captions


def test_healthy_database_raises_no_override_warning(monkeypatch, tmp_path):
    at = _drive(monkeypatch, tmp_path / "ok.db")
    assert not at.exception, at.exception
    warnings = " ".join(str(w.value) for w in at.warning)
    assert "override database could not be read" not in warnings.lower()


def test_invalid_ktc_id_is_flagged_in_planning_mode(monkeypatch, tmp_path):
    at = _drive(monkeypatch, tmp_path / "ok.db", ktc_id="19")
    assert not at.exception, at.exception
    warnings = " ".join(str(w.value) for w in at.warning)
    assert "KTC-ID" in warnings and "3 digits" in warnings


def test_workbook_reports_an_article_setup_overflow(monkeypatch):
    """v34.52: the Result numbering fits, but the Article setup's KTC
    successors continue the same dimension group past 99 variants. The sheet
    used to vanish without a word; the export notes now say why."""
    from engine.workbook import build_result_workbook

    kan = _frame(n=60, system="Kanban")
    ktc = _frame(n=60, system="KTC")
    ktc["Code"] = [f"K{i}" for i in range(60)]
    work = pd.concat([kan, ktc], ignore_index=True)
    work["SystemCategory"] = ["Kanban"] * 60 + ["KTC"] * 60
    work["Monthly_packs"] = 1.0
    work["Spirals_needed"] = [0] * 60 + [1] * 60
    work["Carousel_stockpiles"] = 0
    work["CabinetType"] = ["Kanban"] * 60 + ["Helix"] * 60
    empty = pd.DataFrame()
    ov = {"rows_touched": 0, "unmatched_overrides": [],
          "invalid_overrides": [], "applied_overrides": []}
    raw, problems, _p, _c, notes = build_result_workbook(
        work=work, df_summary=empty, df_audit=empty, df_bucket_compare=empty,
        presentation_compact_df=empty, presentation_detail_df=empty,
        dist_cat_rows=empty, dist_cat_vol=empty, dist_cabtype=empty,
        dist_system=empty, include_planogram=False, include_technical=False,
        ktc_id="191", apply_overrides_ui=False, enable_bulk_routing=False,
        multiple_listings=False, consumption_period_months=12.0,
        content_key="t", _df_run_meta=empty, _bucket_plans=[],
        _listings_arg=None, _listings_in_data=["Tools"], _base_info=None,
        _vend_stats={}, _override_stats=ov,
    )
    from openpyxl import load_workbook
    sheets = load_workbook(BytesIO(raw)).sheetnames
    assert "Article setup" not in sheets
    assert "Article setup sheet was left out" in notes
    assert "99" in notes
