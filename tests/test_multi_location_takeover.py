"""Takeover sheets for articles on several machines (v34.63, owner decision D5).

An article planned in several supply points (``Location_Shared``) has one
stock pool (``Stock_Article_pcs``: its stock in the source list, each source
row counted once). The takeover fills its supply points in order with whole
packs up to each maximum; the rest stays at the HLO on the first supply
point's line and the other lines show 0 at the HLO. Stock is never counted
twice across the sheets. Replicate runs keep their behaviour.
"""
import pandas as pd
import pytest

from engine.fixed_config import STATUS_NOT_PLACED
from engine.multi_location import ARTICLE_STOCK_COL, SHARED_COL
from engine.takeover import (
    HINWEIS_COL,
    STOCK_COL,
    TAKEOVER_COLUMNS,
    build_takeover_frames,
    build_takeover_workbook,
    takeover_summary,
)
from tests.test_takeover import _frame, _row


def _shared(code, sp, pool, *, share=None, **kw):
    return _row(code, SupplyPoint=sp, **{STOCK_COL: pool / 2 if share is None else share,
                                        ARTICLE_STOCK_COL: float(pool), SHARED_COL: True}, **kw)


def _single(code, sp, stock, **kw):
    return _row(code, SupplyPoint=sp, **{STOCK_COL: float(stock), ARTICLE_STOCK_COL: float(stock),
                                        SHARED_COL: False}, **kw)


def _line(frames, sp, code):
    f = dict(frames)[sp]
    return f[f["Kunden Art. Nr."] == code].iloc[0]


def _total(frames, code):
    return sum(float(f.loc[f["Kunden Art. Nr."] == code, "Total"].sum()) for _, f in frames)


def test_the_pool_fills_the_supply_points_in_order_and_the_rest_stays_on_the_first_line():
    # pool 50, VPE 5, maximum 20 pieces (4 compartments x 5) at SP 1 and SP 2
    df = _frame(_shared("A01", 1, 50, PackUnits=5.0, Carousel_stockpiles=4),
                _shared("A01", 2, 50, PackUnits=5.0, Carousel_stockpiles=4))
    frames = build_takeover_frames(df)
    sp1, sp2 = _line(frames, 1, "A01"), _line(frames, 2, "A01")
    assert (sp1["im KTC"], sp1["am HLO"], sp1["Total"]) == (20, 10, 30)
    assert (sp2["im KTC"], sp2["am HLO"], sp2["Total"]) == (20, 0, 20)
    assert _total(frames, "A01") == 50


def test_a_pool_smaller_than_the_first_maximum_leaves_the_second_empty():
    df = _frame(_shared("A01", 1, 15, PackUnits=5.0, Carousel_stockpiles=4),
                _shared("A01", 2, 15, PackUnits=5.0, Carousel_stockpiles=4))
    frames = build_takeover_frames(df)
    assert (_line(frames, 1, "A01")["im KTC"], _line(frames, 1, "A01")["am HLO"]) == (15, 0)
    assert (_line(frames, 2, "A01")["im KTC"], _line(frames, 2, "A01")["am HLO"]) == (0, 0)


def test_only_whole_packs_go_into_the_machines():
    df = _frame(_shared("A01", 1, 23, PackUnits=5.0, Carousel_stockpiles=2),
                _shared("A01", 2, 23, PackUnits=5.0, Carousel_stockpiles=4))
    frames = build_takeover_frames(df)
    assert _line(frames, 1, "A01")["im KTC"] == 10
    assert _line(frames, 2, "A01")["im KTC"] == 10
    assert _line(frames, 1, "A01")["am HLO"] == 3
    assert _total(frames, "A01") == 23


def test_a_shared_kanban_article_stays_at_the_hlo_on_the_first_line():
    kw = dict(SystemCategory="Kanban", CabinetType="Kanban", Carousel_stockpiles=0)
    df = _frame(_shared("K01", 1, 40, **kw), _shared("K01", 2, 40, **kw))
    frames = build_takeover_frames(df)
    assert (_line(frames, 1, "K01")["im KTC"], _line(frames, 1, "K01")["am HLO"]) == (0, 40)
    assert (_line(frames, 2, "K01")["im KTC"], _line(frames, 2, "K01")["am HLO"]) == (0, 0)
    assert _line(frames, 2, "K01")["Schranktyp"] == "Kanban"


def test_a_regrind_helix_article_keeps_its_spiral_for_reground_pieces():
    kw = dict(CabinetType="Helix", Spirals_needed=2, Spiral_capacity=22, Carousel_stockpiles=0,
              Regrind=True)
    df = _frame(_shared("H01", 1, 100, **kw), _shared("H01", 2, 100, **kw))
    frames = build_takeover_frames(df)
    assert _line(frames, 1, "H01")["im KTC"] == 22
    assert _line(frames, 2, "H01")["im KTC"] == 22
    assert _line(frames, 1, "H01")["am HLO"] == 56


def test_a_shared_article_not_placed_in_one_supply_point():
    df = _frame(_shared("A01", 1, 30, Carousel_stockpiles=5, Placement_Status=STATUS_NOT_PLACED),
                _shared("A01", 2, 30, Carousel_stockpiles=5))
    frames = build_takeover_frames(df)
    sp1, sp2 = _line(frames, 1, "A01"), _line(frames, 2, "A01")
    assert (sp1["im KTC"], sp1["am HLO"], sp1["Schranktyp"]) == (0, 25, "Not placed")
    assert (sp2["im KTC"], sp2["am HLO"]) == (5, 0)
    assert _total(frames, "A01") == 30


def test_three_supply_points_and_the_note_names_the_others():
    df = _frame(*[_shared("A01", sp, 9, Carousel_stockpiles=3) for sp in (1, 2, 3)])
    frames = build_takeover_frames(df)
    assert [_line(frames, sp, "A01")["im KTC"] for sp in (1, 2, 3)] == [3, 3, 3]
    assert _line(frames, 1, "A01")[HINWEIS_COL] == "Bestand geteilt mit SP 2, SP 3"
    assert _line(frames, 2, "A01")[HINWEIS_COL] == "Bestand geteilt mit SP 1, SP 3"


def test_the_note_is_only_on_pooled_lines():
    df = _frame(_shared("A01", 1, 10), _shared("A01", 2, 10), _single("B02", 1, 4))
    frames = build_takeover_frames(df)
    for _sp, f in frames:
        assert list(f.columns) == list(TAKEOVER_COLUMNS) + [HINWEIS_COL]
    assert _line(frames, 1, "A01")[HINWEIS_COL] == "Bestand geteilt mit SP 2"
    assert _line(frames, 1, "B02")[HINWEIS_COL] == ""
    for ch in (chr(0x2013), chr(0x2014), chr(0x00B7)):
        assert ch not in _line(frames, 1, "A01")[HINWEIS_COL]


def test_no_note_column_without_a_shared_article():
    df = _frame(_single("A01", 1, 4), _single("B02", 2, 3))
    for _sp, f in build_takeover_frames(df):
        assert list(f.columns) == list(TAKEOVER_COLUMNS)


def test_an_article_in_one_supply_point_keeps_its_stock():
    df = _frame(_shared("A01", 1, 10), _shared("A01", 2, 10),
                _single("B02", 1, 7, Carousel_stockpiles=3))
    line = _line(build_takeover_frames(df), 1, "B02")
    assert (line["im KTC"], line["am HLO"], line["Total"]) == (3, 4, 7)


def test_a_row_that_lost_a_share_to_an_unplanned_machine_keeps_its_whole_stock():
    # AB-100 + AB-102 with AB-102 not planned: one copy, half the stock share,
    # but the pool (the row's whole stock) is taken over.
    df = _frame(_row("C03", SupplyPoint=1, Carousel_stockpiles=3,
                     **{STOCK_COL: 5.0, ARTICLE_STOCK_COL: 10.0, SHARED_COL: False}))
    line = _line(build_takeover_frames(df), 1, "C03")
    assert (line["im KTC"], line["am HLO"], line["Total"]) == (3, 7, 10)


def test_the_pool_is_used_without_the_replicate_switch():
    df = _frame(_shared("A01", 1, 10, Carousel_stockpiles=3),
                _shared("A01", 2, 10, Carousel_stockpiles=3))
    assert sum(r["total"] for r in takeover_summary(build_takeover_frames(df))) == 10


def test_every_copy_shows_the_article_number_once_per_sheet():
    df = _frame(_shared("A01", 1, 4, Kromi_Art_No="140100001000"),
                _shared("A01", 2, 4, Kromi_Art_No="140100001000"))
    frames = build_takeover_frames(df)
    for _sp, f in frames:
        assert list(f["Cust. Prop. Art. Nr."]) == ["140100001000"]


# ---- Replicate mode is unchanged ---------------------------------------------------------

def test_replicate_output_is_unchanged():
    df = _frame(_row("A01", SupplyPoint=1, Carousel_stockpiles=3, **{STOCK_COL: 10.0}),
                _row("A01", SupplyPoint=2, Carousel_stockpiles=3, **{STOCK_COL: 10.0}),
                _row("B02", SupplyPoint=1, PackUnits=4.0, Carousel_stockpiles=1,
                     **{STOCK_COL: 9.0}),
                _row("B02", SupplyPoint=2, PackUnits=4.0, Carousel_stockpiles=1,
                     **{STOCK_COL: 9.0}))
    frames = dict(build_takeover_frames(df, shared_stock=True))
    expected = {
        1: [["A01", 1, 3, 4, 7], ["B02", 4, 4, 1, 5]],
        2: [["A01", 1, 3, 0, 3], ["B02", 4, 4, 0, 4]],
    }
    for sp, rows in expected.items():
        f = frames[sp]
        assert list(f.columns) == list(TAKEOVER_COLUMNS)
        got = f[["Kunden Art. Nr.", "VPE", "im KTC", "am HLO", "Total"]].values.tolist()
        assert got == rows


def test_the_workbook_widens_the_note_column():
    from io import BytesIO

    from openpyxl import load_workbook

    df = _frame(_shared("A01", 1, 10), _shared("A01", 2, 10))
    wb = load_workbook(BytesIO(build_takeover_workbook(build_takeover_frames(df))))
    ws = wb["Takeover sheet SP 1"]
    assert ws.cell(row=1, column=10).value == HINWEIS_COL
    assert ws.column_dimensions["J"].width >= 20


@pytest.mark.parametrize("pool", [0, 1, 7, 49, 50, 51, 1000])
def test_totals_equal_the_pool(pool):
    df = _frame(_shared("A01", 1, pool, PackUnits=5.0, Carousel_stockpiles=4),
                _shared("A01", 2, pool, PackUnits=5.0, Carousel_stockpiles=4))
    frames = build_takeover_frames(df)
    assert _total(frames, "A01") == pool
    for _sp, f in frames:
        assert (pd.to_numeric(f["am HLO"]) >= 0).all()
