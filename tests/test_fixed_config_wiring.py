"""Fixed-configuration mode wiring (v34.52): fingerprint, restore, exports, page.

* The run fingerprint moves when a machine count, the headroom or the move
  toggle changes, but only in fixed mode; every other mode keeps its golden
  fingerprint (the settings fold into the mode-specific position).
* Load & recompute restores the fixed-configuration controls.
* The workbook carries the capacity table and the not-placed list; the
  machine sheets and the planogram hold only placed articles; other modes
  gain nothing.
* The page shows only the controls that apply in fixed mode and keeps the
  hidden values for the other modes.
"""
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from engine.fixed_config import FIXED_MODE, STATUS_NOT_PLACED
from engine.run_fingerprint import FingerprintInputs, compute_run_fingerprint
from engine.run_restore import CONTROL_KEYS, build_seed, capture_ui_state
from tests._paths import PLANNER_PAGE
from tests.test_run_fingerprint import _BASE, GOLDEN

# ---- fingerprint -------------------------------------------------------------------

_FIXED = (((1, 2, 1, 0, 0, 0), (2, 1, 1, 0, 0, 0)), 10.0, True)


def _fp(**kw):
    d = dict(_BASE)
    d.update(kw)
    return compute_run_fingerprint(FingerprintInputs(**d))


def test_fixed_settings_do_not_touch_other_modes():
    assert _fp(fixed_config=_FIXED) == GOLDEN          # Capped baseline
    assert _fp(op_mode="", fixed_config=_FIXED) == _fp(op_mode="")


def test_fixed_settings_move_the_fingerprint_in_fixed_mode():
    base = _fp(op_mode=FIXED_MODE, fixed_config=_FIXED)
    assert len(base) == len(GOLDEN)
    for alt in (
        (((1, 3, 1, 0, 0, 0), (2, 1, 1, 0, 0, 0)), 10.0, True),
        (((1, 2, 1, 0, 0, 0), (2, 1, 1, 0, 0, 1)), 10.0, True),
        (_FIXED[0], 15.0, True),
        (_FIXED[0], 10.0, False),
    ):
        assert _fp(op_mode=FIXED_MODE, fixed_config=alt) != base, alt


# ---- restore -----------------------------------------------------------------------

def test_fixed_controls_are_restorable():
    keys = set(CONTROL_KEYS)
    assert {"ks_fixed_headroom", "ks_fixed_spill"} <= keys
    for sp in range(1, 11):
        for t in ("helix", "carousel", "locker_a", "locker_b", "locker_c"):
            assert f"ks_fixed_{t}_{sp}" in keys
    state = {"ks_fixed_headroom": 15.0, "ks_fixed_spill": False,
             "ks_fixed_helix_2": 3, "ks_op_mode": "Fixed configuration (existing machines)"}
    snap = capture_ui_state(state)
    assert build_seed(snap, ["A"]) == state


# ---- workbook ------------------------------------------------------------------------

def _workbook(res):
    from engine.workbook import build_result_workbook
    empty = pd.DataFrame()
    ov = {"rows_touched": 0, "unmatched_overrides": [], "invalid_overrides": [],
          "applied_overrides": []}
    raw, problems, _p, _c, _notes = build_result_workbook(
        work=res.work, df_summary=empty, df_audit=empty, df_bucket_compare=empty,
        presentation_compact_df=empty, presentation_detail_df=empty,
        dist_cat_rows=empty, dist_cat_vol=empty, dist_cabtype=empty,
        dist_system=empty, include_planogram=True, include_technical=False,
        ktc_id="191", apply_overrides_ui=False, enable_bulk_routing=False,
        multiple_listings=False, consumption_period_months=16.0,
        content_key="t", _df_run_meta=empty, _bucket_plans=res.bucket_plans,
        _listings_arg=None, _listings_in_data=["Tools"], _base_info=None,
        _vend_stats={}, _override_stats=ov,
    )
    return load_workbook(BytesIO(raw)), problems


def _sheet_frame(wb, name):
    rows = list(wb[name].values)
    return pd.DataFrame(rows[1:], columns=rows[0])


def _squeezed():
    """A fixed run where the machines are too small for every article."""
    from engine.plan import run_plan
    from tests.test_run_plan_equivalence import _boundary_frame, _params
    df = _boundary_frame()
    p = _params(df, op_mode=FIXED_MODE, fixed_machines=((1, 1, 0, 0, 0, 0),),
                fixed_headroom_pct=90.0, fixed_allow_spill=False)
    return run_plan(df, pd.DataFrame(), p)


def test_fixed_workbook_has_capacity_and_not_placed_sheets():
    res = _squeezed()
    n_not = int((res.work["Placement_Status"] == STATUS_NOT_PLACED).sum())
    assert n_not > 0
    wb, problems = _workbook(res)
    assert problems == [], problems
    cap = _sheet_frame(wb, "Fixed_Configuration")
    assert list(cap["Machine"]) == ["Helix"]
    assert int(cap.loc[0, "Usable"]) == 7 and int(cap.loc[0, "Machines"]) == 1
    miss = _sheet_frame(wb, "Not_placed")
    assert len(miss) == n_not
    assert "Placement_Note" in miss.columns
    result = _sheet_frame(wb, "Result")
    assert {"Placement_Rank", "Placement_Status", "Placement_Note"} <= set(result.columns)


def test_machine_sheets_hold_only_placed_articles():
    res = _squeezed()
    wb, _ = _workbook(res)
    placed_codes = set(res.work.loc[
        (res.work["Placement_Status"] != STATUS_NOT_PLACED)
        & (res.work["SystemCategory"] == "KTC"), "Code"].astype(str))
    for sheet in ("Helix_only", "Carousel_only"):
        frame = _sheet_frame(wb, sheet)
        if len(frame):
            assert set(frame["Code"].astype(str)) <= placed_codes, sheet


def test_standard_workbook_gains_no_fixed_sheets():
    from engine.plan import run_plan
    from tests.test_run_plan_equivalence import _boundary_frame, _params
    df = _boundary_frame()
    res = run_plan(df, pd.DataFrame(), _params(df))
    wb, _ = _workbook(res)
    assert "Fixed_Configuration" not in wb.sheetnames
    assert "Not_placed" not in wb.sheetnames


def test_run_meta_rows_describe_the_configuration():
    from engine.fixed_config import run_meta_rows
    rows = {r["Key"]: r["Value"] for r in run_meta_rows(_squeezed().bucket_plans)}
    assert rows["Fixed configuration machines"] == "All: 1 Helix"
    assert rows["Fixed configuration headroom (%)"] == 90.0
    assert rows["Move overflow to another machine type"] == "off"
    assert rows["Articles not placed (no space)"] > 0


# ---- page ---------------------------------------------------------------------------

FIXED_LABEL = "Fixed configuration (existing machines)"
_HIDDEN_IN_FIXED = ("Carousel fill ceiling", "Consolidate underused cabinets",
                    "Empty-cabinet threshold (%)", "Capacity buffer (%)")


def _xlsx():
    df = pd.DataFrame({
        "ItemCode": [f"V{i:03d}" for i in range(8)],
        "ItemName": [f"Bohrer D{i + 3},5" for i in range(8)],
        "Cons16": [(i + 1) * 600.0 for i in range(8)],
    })
    bio = BytesIO()
    with pd.ExcelWriter(bio, engine="openpyxl") as w:
        df.to_excel(w, index=False, sheet_name="Sheet1")
    return bio.getvalue(), list(df.columns)


def _drive(monkeypatch, tmp_path, *, extra_state=None, run=True):
    import streamlit as st
    from streamlit.testing.v1 import AppTest
    monkeypatch.setenv("KROMI_DB_PATH", str(tmp_path / "fx.db"))
    st.cache_data.clear()
    raw, cols = _xlsx()
    seed = build_seed({"cm_code": "ItemCode", "cm_desc1": "ItemName",
                       "cm_cons": "Cons16", "ks_sheet_tools": "Sheet1"}, cols)
    seed.update({"ks_hide_optional": False, "ks_ai_colmap": False,
                 "_std_special_mapped": False})
    seed.update(extra_state or {})
    at = AppTest.from_file(PLANNER_PAGE, default_timeout=300)
    at.session_state["_reload_ctx"] = {
        "run_id": 1, "bytes": raw, "filename": "input.xlsx",
        "customer": "FxCustomer", "site": "FxSite", "classifications": {},
        "override_set_id": None}
    at.session_state["_pending_restore"] = seed
    at.run()
    if run:
        at.session_state["_force_run"] = True
        at.run()
    assert not at.exception, at.exception
    return at


def _labels(at):
    return {str(w.label) for w in list(at.number_input) + list(at.checkbox) + list(at.radio)}


def test_fixed_mode_shows_only_relevant_controls(monkeypatch, tmp_path):
    at = _drive(monkeypatch, tmp_path, run=False, extra_state={"ks_op_mode": FIXED_LABEL})
    labels = _labels(at)
    for hidden in _HIDDEN_IN_FIXED:
        assert hidden not in labels, hidden
    assert not any(lbl.startswith("Tools + PPE handling") for lbl in labels)
    assert "Headroom to keep free (%)" in labels
    assert any(lbl.startswith("Move overflow") for lbl in labels)
    keys = {getattr(w, "key", "") for w in at.number_input}
    assert {"ks_fixed_helix_1", "ks_fixed_carousel_1", "ks_fixed_locker_a_1"} <= keys
    assert "ks_max_carousels" not in keys


def test_standard_mode_keeps_its_controls(monkeypatch, tmp_path):
    at = _drive(monkeypatch, tmp_path, run=False)
    labels = _labels(at)
    for shown in _HIDDEN_IN_FIXED:
        assert shown in labels, shown
    assert "Headroom to keep free (%)" not in labels


def test_hidden_values_survive_a_mode_round_trip(monkeypatch, tmp_path):
    at = _drive(monkeypatch, tmp_path, run=False, extra_state={"ks_buffer": 35.0})
    next(s for s in at.selectbox if s.key == "ks_op_mode").select(FIXED_LABEL).run()
    assert not at.exception, at.exception
    next(s for s in at.selectbox if s.key == "ks_op_mode").select(
        "Standard (best fit per tool)").run()
    assert not at.exception, at.exception
    buf = next(n for n in at.number_input if n.key == "ks_buffer")
    assert buf.value == 35.0


def test_fixed_run_reports_the_fit(monkeypatch, tmp_path):
    at = _drive(monkeypatch, tmp_path, extra_state={
        "ks_op_mode": FIXED_LABEL, "ks_fixed_helix_1": 1, "ks_fixed_carousel_1": 0,
        "ks_fixed_headroom": 90.0, "ks_fixed_spill": False})
    text = " ".join(str(m.value) for m in list(at.markdown) + list(at.info)
                    + list(at.warning) + list(at.caption))
    assert "Operational mode: fixed configuration" in text
    assert "not placed" in text.lower()
    subheaders = [str(s.value) for s in at.subheader]
    assert "Fit into the configured machines" in subheaders
    assert "Configured machines (KTC only)" in subheaders


def test_no_machines_stops_with_a_clear_message(monkeypatch, tmp_path):
    at = _drive(monkeypatch, tmp_path, run=False, extra_state={
        "ks_op_mode": FIXED_LABEL, "ks_fixed_helix_1": 0, "ks_fixed_carousel_1": 0})
    errors = " ".join(str(e.value) for e in at.error)
    assert "at least one machine" in errors


def _xlsx_with_program():
    df = pd.DataFrame({
        "ItemCode": [f"V{i:03d}" for i in range(8)],
        "ItemName": [f"Bohrer D{i + 3},5" for i in range(8)],
        "Cons16": [(i + 1) * 600.0 for i in range(8)],
        "Area": ["KTC-A" if i % 2 else "KTC-B" for i in range(8)],
    })
    bio = BytesIO()
    with pd.ExcelWriter(bio, engine="openpyxl") as w:
        df.to_excel(w, index=False, sheet_name="Sheet1")
    return bio.getvalue(), list(df.columns)


def test_program_mapping_places_each_location_in_its_own_machines(monkeypatch, tmp_path):
    import streamlit as st
    from streamlit.testing.v1 import AppTest
    monkeypatch.setenv("KROMI_DB_PATH", str(tmp_path / "fx.db"))
    st.cache_data.clear()
    raw, cols = _xlsx_with_program()
    seed = build_seed({"cm_code": "ItemCode", "cm_desc1": "ItemName",
                       "cm_cons": "Cons16", "ks_sheet_tools": "Sheet1"}, cols)
    seed.update({"ks_hide_optional": False, "ks_ai_colmap": False,
                 "_std_special_mapped": False, "ks_op_mode": FIXED_LABEL,
                 "ks_n_sp": 2, "ks_fixed_helix_2": 0})
    at = AppTest.from_file(PLANNER_PAGE, default_timeout=300)
    at.session_state["_reload_ctx"] = {
        "run_id": 1, "bytes": raw, "filename": "input.xlsx",
        "customer": "FxCustomer", "site": "FxSite", "classifications": {},
        "override_set_id": None}
    at.session_state["_pending_restore"] = seed
    at.run()
    next(s for s in at.selectbox if s.key == "cm_program").select("Area").run()
    next(s for s in at.selectbox if s.key == "prog_sp::KTC-B").select(2).run()
    at.session_state["_force_run"] = True
    at.run()
    assert not at.exception, at.exception
    captions = " ".join(str(c.value) for c in at.caption)
    assert "items assigned by the Program column" in captions
    assert "every item replicated" not in captions
    assert "has been duplicated" not in captions
    frames = [d.value for d in at.dataframe]
    cap = next(f for f in frames if "Free (usable)" in getattr(f, "columns", []))
    assert set(cap["Supply point"]) == {"SP 1 (KTC-A)", "SP 2 (KTC-B)"}
    sp2 = cap[cap["Supply point"] == "SP 2 (KTC-B)"]
    assert list(sp2["Machine"]) == ["Carousel"]          # SP 2 has no Helix
