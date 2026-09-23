"""Takeover sheets wired through the planner (v34.56).

* The stock column is an optional mapping; the planning base sums it per
  article and keeps the raw text of a mapped category column (the KDS
  template's Bezeichnung 1) for the sheet.
* The plan carries both columns through; the result workbook adds one
  takeover sheet per supply point, and none when no stock column is mapped
  (every existing export stays as it was).
* The takeover numbers are the customer-property numbers of the Result and
  Article setup sheets.
* Load & recompute restores the stock mapping, and mapping it changes the
  run fingerprint (so results are recomputed) only when it is mapped.
"""
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from engine.fixed_config import FIXED_MODE
from engine.preprocessing import prepare_planning_base
from engine.takeover import CATEGORY_TEXT_COL, STOCK_COL, TAKEOVER_COLUMNS
from tests._paths import PLANNER_PAGE


# ---- planning base -------------------------------------------------------------------

def _base(**extra):
    df = pd.DataFrame({
        "Listing": ["Tools"] * 3, "Code": ["A", "A", "B"], "SupplierCode": ["", "", ""],
        "Description": ["d a", "d a", "d b"], "Description_2": ["", "", ""],
        "Consumption_pcs": [10.0, 5.0, 3.0], "PackUnits": [1.0, 1.0, 1.0],
        "ProductCategory": ["drills", "drills", "mills"], "SizeCategory": ["", "", ""],
    })
    for k, v in extra.items():
        df[k] = v
    return df


def test_dedup_sums_the_stock_and_keeps_the_category_text():
    df = _base(**{STOCK_COL: [3.0, 4.0, 1.0],
                  CATEGORY_TEXT_COL: ["", "Stufenbohrer VHM", "Schaftfräser VHM"]})
    out, _ = prepare_planning_base(df, dedup_mode="code", year_mode="all", has_year=False)
    got = out.set_index("Code")
    assert got.loc["A", STOCK_COL] == 7.0 and got.loc["B", STOCK_COL] == 1.0
    assert got.loc["A", CATEGORY_TEXT_COL] == "Stufenbohrer VHM"


def test_without_the_new_columns_the_planning_base_is_unchanged():
    out, _ = prepare_planning_base(_base(), dedup_mode="code", year_mode="all",
                                   has_year=False)
    assert STOCK_COL not in out.columns and CATEGORY_TEXT_COL not in out.columns


# ---- plan and workbook -------------------------------------------------------------------

def _run(n_sp=2, sp_mode="partition", stock=True, **kw):
    from engine.plan import run_plan
    from tests.test_run_plan_equivalence import _boundary_frame, _params
    df = _boundary_frame()
    if "ToolClass" not in df.columns:
        df["ToolClass"] = "other"          # numbers need a tool class
    if stock:
        df[STOCK_COL] = [float(i * 13 % 40) for i in range(len(df))]
    p = _params(df, n_supply_points=n_sp, sp_mode=sp_mode, **kw)
    return run_plan(df, pd.DataFrame(), p)


def _workbook(res, **extra):
    from engine.workbook import build_result_workbook
    empty = pd.DataFrame()
    ov = {"rows_touched": 0, "unmatched_overrides": [], "invalid_overrides": [],
          "applied_overrides": []}
    raw, problems, _p, _c, notes = build_result_workbook(
        work=res.work, df_summary=empty, df_audit=empty, df_bucket_compare=empty,
        presentation_compact_df=empty, presentation_detail_df=empty,
        dist_cat_rows=empty, dist_cat_vol=empty, dist_cabtype=empty,
        dist_system=empty, include_planogram=False, include_technical=False,
        ktc_id="191", apply_overrides_ui=False, enable_bulk_routing=False,
        multiple_listings=False, consumption_period_months=16.0,
        content_key="t", _df_run_meta=empty, _bucket_plans=res.bucket_plans,
        _listings_arg=None, _listings_in_data=["Tools"], _base_info=None,
        _vend_stats={}, _override_stats=ov, **extra,
    )
    assert problems == [], problems
    return load_workbook(BytesIO(raw))


def _sheet(wb, name):
    rows = list(wb[name].values)
    return pd.DataFrame(rows[1:], columns=rows[0])


def test_the_plan_carries_the_stock_through():
    res = _run()
    assert STOCK_COL in res.work.columns
    assert res.work[STOCK_COL].sum() == sum(float(i * 13 % 40) for i in range(14))


def test_one_takeover_sheet_per_supply_point():
    res = _run()
    wb = _workbook(res)
    names = [s for s in wb.sheetnames if s.startswith("Takeover sheet")]
    assert names == ["Takeover sheet SP 1", "Takeover sheet SP 2"]
    frames = [_sheet(wb, n) for n in names]
    for f in frames:
        assert list(f.columns) == list(TAKEOVER_COLUMNS)
    assert sum(len(f) for f in frames) == len(res.work)
    assert sum(float(f["Total"].sum()) for f in frames) == float(res.work[STOCK_COL].sum())
    # The numbers are the Article setup's customer-property numbers.
    setup = _sheet(wb, "Article setup")
    cust = dict(zip(setup.loc[setup["Property"] == "Customer property", "Customer article No"],
                    setup.loc[setup["Property"] == "Customer property", "Kromi_Art_No"]))
    for f in frames:
        for code, number in zip(f["Kunden Art. Nr."], f["Cust. Prop. Art. Nr."]):
            assert cust[str(code)] == str(number)


def test_no_stock_mapping_means_no_takeover_sheet():
    wb = _workbook(_run(stock=False))
    assert not [s for s in wb.sheetnames if s.startswith("Takeover")]


def test_replicated_stock_is_counted_once():
    res = _run(sp_mode="replicate")
    total = float(res.work.drop_duplicates("Code")[STOCK_COL].sum())
    wb = _workbook(res, _takeover_shared_stock=True)
    frames = [_sheet(wb, s) for s in wb.sheetnames if s.startswith("Takeover")]
    assert len(frames) == 2
    assert sum(float(f["Total"].sum()) for f in frames) == total


def test_fixed_configuration_takeover_leaves_not_placed_articles_at_hlo():
    res = _run(n_sp=1, op_mode=FIXED_MODE, fixed_machines=((1, 1, 0, 0, 0, 0),),
               fixed_headroom_pct=90.0, fixed_allow_spill=False)
    wb = _workbook(res)
    t = _sheet(wb, "Takeover sheet SP 1")
    miss = t[t["Schranktyp"] == "Not placed"]
    assert len(miss) > 0
    assert (miss["im KTC"] == 0).all() and (miss["am HLO"] == miss["Total"]).all()


# ---- restore and page ------------------------------------------------------------------------

def test_stock_mapping_is_restorable():
    from engine.run_restore import MAPPING_KEYS, REQUIRED_MAPPING_KEYS
    assert "cm_stock" in MAPPING_KEYS and "cm_stock" not in REQUIRED_MAPPING_KEYS


def _kds_like_bytes():
    df = pd.DataFrame({
        "Kunden Artikelnummer / Customer Art.-No": ["60B0001", "60B0002", "60F0003"],
        "Bezeichnung 1 / Description 1": ["Stufenbohrer VHM", "Spiralbohrer VHM",
                                          "Schaftfräser VHM"],
        "Bezeichnung 2 / Description 2": ["Drm. 4,80/ 8,00", "Drm. 3,50 GL61",
                                          "D 12 Z4"],
        "Verbrauch 12 Monate": [120.0, 600.0, 24.0],
        "Preis- basis Stück / Price basis pcs.": [1, 10, 1],
        "AKTUELLER BESTAND 31.08.2026": [9, 310, 0],
        "Mindestbestand / Min. stock": [1, 1, 1],
    })
    bio = BytesIO()
    with pd.ExcelWriter(bio, engine="openpyxl") as w:
        df.to_excel(w, index=False, sheet_name="Tabelle1")
    return bio.getvalue(), list(df.columns)


def _drive(monkeypatch, tmp_path, *, map_stock=True):
    import streamlit as st
    from streamlit.testing.v1 import AppTest
    import engine.workbook as wbmod
    from engine.run_restore import NOT_AVAIL, build_seed

    captured = {}
    orig = wbmod.build_takeover_frames

    def capture(work, **kw):
        frames = orig(work, **kw)
        captured["frames"] = frames
        return frames
    monkeypatch.setattr(wbmod, "build_takeover_frames", capture)
    monkeypatch.setenv("KROMI_DB_PATH", str(tmp_path / "tk.db"))
    st.cache_data.clear()
    raw, cols = _kds_like_bytes()
    seed = build_seed({"cm_code": cols[0], "cm_desc1": cols[2], "cm_cons": cols[3],
                       "ks_sheet_tools": "Tabelle1"}, cols)
    seed.update({"ks_hide_optional": False, "ks_ai_colmap": False,
                 "_std_special_mapped": False})
    at = AppTest.from_file(PLANNER_PAGE, default_timeout=300)
    at.session_state["_reload_ctx"] = {
        "run_id": 0, "bytes": raw, "filename": "kds.xlsx", "customer": "TkC",
        "site": "TkS", "classifications": {}, "override_mode": "none",
        "override_set_id": None}
    at.session_state["_pending_restore"] = seed
    at.run()
    at.selectbox(key="cm_prod").select(cols[1])
    at.selectbox(key="cm_pack").select(cols[4])
    at.run()
    if map_stock:
        assert at.selectbox(key="cm_stock").value == cols[5]    # found automatically
    else:
        at.selectbox(key="cm_stock").select(NOT_AVAIL)
    at.run()
    at.session_state["_force_run"] = True
    at.run()
    assert not at.exception, at.exception
    btn = next(b for b in at.button if getattr(b, "key", "") == "btn_prepare_exports")
    btn.click()
    at.run()
    assert not at.exception, at.exception
    return at, captured


def test_page_builds_the_takeover_from_the_mapped_stock(monkeypatch, tmp_path):
    at, captured = _drive(monkeypatch, tmp_path)
    frames = dict(captured["frames"])
    t = frames[1].set_index("Kunden Art. Nr.")
    assert t.loc["60B0001", "Bezeichnung 1"] == "Stufenbohrer VHM"
    assert t.loc["60B0001", "Bezeichnung 2"] == "Drm. 4,80/ 8,00"
    assert int(t.loc["60B0002", "VPE"]) == 10
    assert int(t.loc["60B0002", "Total"]) == 310
    assert int(t["Total"].sum()) == 319
    captions = " ".join(str(c.value) for c in at.caption)
    assert "Takeover sheets" in captions
    fp_with = at.session_state["_run_fingerprint"]

    at2, captured2 = _drive(monkeypatch, tmp_path, map_stock=False)
    assert captured2.get("frames", []) == []
    assert at2.session_state["_run_fingerprint"] != fp_with


def test_a_new_file_starts_without_the_previous_stock_mapping():
    from engine.colmap import reset_mapping_state_on_file_change
    state = {"cm_stock": "Bestand", "_mapping_file_id": "old"}
    assert reset_mapping_state_on_file_change(state, "new")
    assert "cm_stock" not in state


def test_result_sheet_shows_the_mapped_stock_next_to_the_pack_size():
    from engine.export_shaping import user_view
    df = pd.DataFrame({"Code": ["A"], "PackUnits": [10.0], STOCK_COL: [45.0],
                       "Consumption_pcs": [1.0]})
    cols = list(user_view(df).columns)
    assert cols.index(STOCK_COL) == cols.index("PackUnits") + 1
