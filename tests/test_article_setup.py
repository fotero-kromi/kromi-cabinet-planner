"""Contracts for the Article setup sheet (v34.43).

Every KTC article receives TWO numbers on the new sheet: the predecessor
(customer property, today's fixed-"10" counter scheme, byte-identical to the
number the Result sheet already shows for that article's first row) and the
successor (KROMI property, the dimension/description scheme Kanban uses).
Kanban articles appear once with their existing Result-sheet number. The
dimension scheme runs one shared variant pool across the WHOLE catalog, so a
successor whose (code, dimension) fingerprint collides with a Kanban article
continues that group's variant counter instead of colliding. The sheet lists
unique articles (by Listing + Code, first occurrence wins), carries the
Replaces / Replaced by linkage, and every number keeps the four invariants
(12 chars, all digits, trailing 0, unique).

v34.45 sheet contract: Kromi_Art_No leads, the article code column is titled
"Customer article No", Listing and CabinetType are dropped, and Kanban rows
are Customer property (Kanban stock is invoiced to the customer at delivery;
only the KTC successor is KROMI property).
"""

from io import BytesIO

import pandas as pd
import pytest

from engine.kromi_numbering import (
    augment_for_export,
    build_article_setup,
)

KTC = "191"


def _frame(rows):
    base_cols = {
        "Code": "", "Listing": "Tools", "Description": "",
        "SystemCategory": "KTC", "CabinetType": "Helix",
        "ToolClass": "solid_carbide_drill", "PackUnits": 1.0,
    }
    recs = []
    for r in rows:
        rec = dict(base_cols)
        rec.update(r)
        recs.append(rec)
    return pd.DataFrame(recs)


def _numbers(df):
    return df["Kromi_Art_No"].tolist()


# ---- shape -----------------------------------------------------------------

def test_ktc_twice_kanban_once():
    df = _frame([
        {"Code": "A1", "Description": "Drill D8,5", "SystemCategory": "KTC"},
        {"Code": "B1", "Description": "Mill D12", "SystemCategory": "Kanban",
         "CabinetType": "Kanban", "ToolClass": "solid_end_mill"},
    ])
    out = build_article_setup(df, KTC)
    assert out is not None
    assert len(out) == 3  # A1 twice, B1 once
    assert list(out["Customer article No"]) == ["A1", "A1", "B1"]


def test_column_set_and_order():
    df = _frame([{"Code": "A1", "Description": "Drill D8"}])
    out = build_article_setup(df, KTC)
    assert list(out.columns) == [
        "Kromi_Art_No", "Customer article No", "Description", "System",
        "Property", "Replaces", "Replaced by", "PackUnits",
    ]


def test_property_labels():
    df = _frame([
        {"Code": "A1", "Description": "Drill D8"},
        {"Code": "B1", "Description": "Mill D12", "SystemCategory": "Kanban",
         "CabinetType": "Kanban"},
    ])
    out = build_article_setup(df, KTC)
    # KTC pair: predecessor customer / successor KROMI. Kanban stock is
    # invoiced to the customer at delivery, so it is Customer property too.
    assert list(out["Property"]) == [
        "Customer property", "KROMI property", "Customer property",
    ]
    assert list(out["System"]) == ["KTC", "KTC", "Kanban"]


# ---- number schemes --------------------------------------------------------

def test_predecessor_matches_result_sheet_number():
    df = _frame([
        {"Code": "A1", "Description": "Drill D8"},
        {"Code": "A2", "Description": "Drill D9"},
    ])
    result_numbers = augment_for_export(df, KTC)["Kromi_Art_No"].tolist()
    out = build_article_setup(df, KTC)
    predecessors = out[out["Property"] == "Customer property"]
    assert list(predecessors["Kromi_Art_No"]) == result_numbers
    # Today's scheme: KTC-ID + "10" + counter + "000"
    assert result_numbers[0] == "191100001000"
    assert result_numbers[1] == "191100002000"


def test_kanban_number_matches_result_sheet_number():
    df = _frame([
        {"Code": "K1", "Description": "Mill D12 lang", "SystemCategory": "Kanban",
         "CabinetType": "Kanban", "ToolClass": "solid_end_mill"},
    ])
    result_number = augment_for_export(df, KTC)["Kromi_Art_No"].iloc[0]
    out = build_article_setup(df, KTC)
    assert out["Kromi_Art_No"].iloc[0] == result_number
    assert result_number.startswith("19116")  # mills -> code 16


def test_successor_uses_dimension_scheme_not_fixed_10():
    df = _frame([{"Code": "A1", "Description": "Drill D8,5x120",
                  "ToolClass": "solid_carbide_drill"}])
    out = build_article_setup(df, KTC)
    succ = out[out["Property"] == "KROMI property"]["Kromi_Art_No"].iloc[0]
    # drills -> code 13; dimension = first 4 digits of description (8,5,1,2)
    assert succ == "191138512000"


def test_successor_holder_anatomy():
    df = _frame([{"Code": "H1", "Description": "HSK63 Aufnahme",
                  "ToolClass": "hsk_holder"}])
    out = build_article_setup(df, KTC)
    succ = out[out["Property"] == "KROMI property"]["Kromi_Art_No"].iloc[0]
    # holder code 20008 + 1 dimension digit + variant + trailing 0
    assert succ == "191200086000"


# ---- shared variant pool ---------------------------------------------------

def test_fingerprint_collision_continues_variant_pool():
    # Kanban article and KTC successor share (code 13, dim 8500): the Kanban
    # row keeps its Result-sheet variant 00; the successor takes 01.
    df = _frame([
        {"Code": "K1", "Description": "Bohrer D8,5", "SystemCategory": "Kanban",
         "CabinetType": "Kanban"},
        {"Code": "A1", "Description": "Bohrer D8,5 VHM"},
    ])
    result_kanban = augment_for_export(df, KTC)
    kanban_no = result_kanban.loc[result_kanban["Code"] == "K1",
                                  "Kromi_Art_No"].iloc[0]
    out = build_article_setup(df, KTC)
    assert out.loc[(out["Customer article No"] == "K1"),
                   "Kromi_Art_No"].iloc[0] == kanban_no
    succ = out[(out["Customer article No"] == "A1")
               & (out["Property"] == "KROMI property")]["Kromi_Art_No"].iloc[0]
    assert kanban_no.endswith("000") and kanban_no[-3:-1] == "00"
    assert succ[-3:-1] == "01"  # continues the pool instead of colliding
    assert succ != kanban_no


def test_successor_variants_continue_after_full_frame_duplicates():
    # Replicate mode: the same Kanban code appears once per SP on the Result
    # sheet, consuming variants 00 and 01 there. The successor sharing the
    # fingerprint must start at 02, so it never collides with any Result-sheet
    # number either.
    df = _frame([
        {"Code": "K1", "Description": "Bohrer D8,5", "SystemCategory": "Kanban",
         "CabinetType": "Kanban", "SupplyPoint": 1},
        {"Code": "K1", "Description": "Bohrer D8,5", "SystemCategory": "Kanban",
         "CabinetType": "Kanban", "SupplyPoint": 2},
        {"Code": "A1", "Description": "Bohrer D8,5 VHM", "SupplyPoint": 1},
        {"Code": "A1", "Description": "Bohrer D8,5 VHM", "SupplyPoint": 2},
    ])
    out = build_article_setup(df, KTC)
    assert len(out) == 3  # K1 once, A1 twice
    succ = out[(out["Customer article No"] == "A1")
               & (out["Property"] == "KROMI property")]["Kromi_Art_No"].iloc[0]
    assert succ[-3:-1] == "02"


def test_unique_by_code_and_predecessor_is_first_occurrence():
    df = _frame([
        {"Code": "A1", "Description": "Drill D8", "SupplyPoint": 1},
        {"Code": "A2", "Description": "Drill D9", "SupplyPoint": 1},
        {"Code": "A1", "Description": "Drill D8", "SupplyPoint": 2},
        {"Code": "A2", "Description": "Drill D9", "SupplyPoint": 2},
    ])
    result_numbers = augment_for_export(df, KTC)["Kromi_Art_No"].tolist()
    out = build_article_setup(df, KTC)
    assert len(out) == 4  # two unique KTC articles x 2 rows
    predecessors = out[out["Property"] == "Customer property"]
    # First occurrences on the Result sheet are rows 0 (A1) and 1 (A2).
    assert list(predecessors["Kromi_Art_No"]) == [result_numbers[0],
                                                  result_numbers[1]]


# ---- linkage ---------------------------------------------------------------

def test_linkage_columns():
    df = _frame([
        {"Code": "A1", "Description": "Drill D8"},
        {"Code": "B1", "Description": "Mill D12", "SystemCategory": "Kanban",
         "CabinetType": "Kanban", "ToolClass": "solid_end_mill"},
    ])
    out = build_article_setup(df, KTC)
    pred = out.iloc[0]; succ = out.iloc[1]; kanban = out.iloc[2]
    assert pred["Replaced by"] == succ["Kromi_Art_No"]
    assert pred["Replaces"] == ""
    assert succ["Replaces"] == pred["Kromi_Art_No"]
    assert succ["Replaced by"] == ""
    assert kanban["Replaces"] == "" and kanban["Replaced by"] == ""


# ---- invariants ------------------------------------------------------------

def test_all_numbers_unique_and_well_formed():
    rows = []
    for i in range(40):
        rows.append({"Code": f"T{i}", "Description": f"Bohrer D8,5 nr {i}"})
    for i in range(15):
        rows.append({"Code": f"N{i}", "Description": f"Bohrer D8,5 kb {i}",
                     "SystemCategory": "Kanban", "CabinetType": "Kanban"})
    out = build_article_setup(_frame(rows), KTC)
    nums = _numbers(out)
    assert len(nums) == 40 * 2 + 15
    assert len(set(nums)) == len(nums)
    for n in nums:
        assert len(n) == 12 and n.isdigit() and n.endswith("0")


def test_variant_overflow_raises():
    # >100 articles sharing one (code, dimension) fingerprint exhausts the
    # 2-digit variant field across Kanban + successors.
    rows = [{"Code": f"T{i}", "Description": "Bohrer D8,5"} for i in range(101)]
    with pytest.raises(ValueError):
        build_article_setup(_frame(rows), KTC)


# ---- graceful omission -----------------------------------------------------

def test_returns_none_without_prerequisites():
    df = _frame([{"Code": "A1", "Description": "Drill D8"}])
    assert build_article_setup(df, "19") is None          # bad KTC-ID
    assert build_article_setup(df, "") is None
    assert build_article_setup(df.drop(columns=["ToolClass"]), KTC) is None
    assert build_article_setup(
        df.drop(columns=["Description"]), KTC) is None    # no description col
    empty = build_article_setup(df.iloc[0:0], KTC)
    assert empty is not None and len(empty) == 0


# ---- workbook integration --------------------------------------------------

def _minimal_workbook_bytes(work, ktc_id):
    from engine.workbook import build_result_workbook

    empty = pd.DataFrame()
    ov_stats = {"rows_touched": 0, "unmatched_overrides": [],
                "invalid_overrides": [], "applied_overrides": []}
    return build_result_workbook(
        work=work, df_summary=empty, df_audit=empty, df_bucket_compare=empty,
        presentation_compact_df=empty, presentation_detail_df=empty,
        dist_cat_rows=empty, dist_cat_vol=empty, dist_cabtype=empty,
        dist_system=empty, include_planogram=False, include_technical=False,
        ktc_id=ktc_id, apply_overrides_ui=False, enable_bulk_routing=False,
        multiple_listings=False, consumption_period_months=12.0,
        content_key="test", _df_run_meta=empty, _bucket_plans=[],
        _listings_arg=None, _listings_in_data=["Tools"], _base_info=None,
        _vend_stats={}, _override_stats=ov_stats,
    )


def _work_for_workbook():
    df = _frame([
        {"Code": "A1", "Description": "Drill D8,5"},
        {"Code": "B1", "Description": "Mill D12", "SystemCategory": "Kanban",
         "CabinetType": "Kanban", "ToolClass": "solid_end_mill"},
    ])
    df["Monthly_packs"] = 1.0
    df["Spirals_needed"] = 1
    df["Carousel_stockpiles"] = 0
    return df


def test_workbook_carries_article_setup_sheet():
    from openpyxl import load_workbook

    xbytes, problems, _p, _c, _n = _minimal_workbook_bytes(_work_for_workbook(), KTC)
    wb = load_workbook(BytesIO(xbytes))
    assert "Article setup" in wb.sheetnames
    ws = wb["Article setup"]
    header = [c.value for c in ws[1]]
    assert header[:6] == ["Kromi_Art_No", "Customer article No", "Description",
                          "System", "Property", "Replaces"]
    assert ws.max_row == 1 + 3  # header + (A1 x2, B1 x1)


def test_workbook_omits_sheet_without_valid_ktc_id():
    from openpyxl import load_workbook

    xbytes, _pr, _p, _c, _n = _minimal_workbook_bytes(_work_for_workbook(), "")
    wb = load_workbook(BytesIO(xbytes))
    assert "Article setup" not in wb.sheetnames
